"""Регрессия на произвольных входных документах, не связанных с контрольной парой.

Запуск: python -m unittest -v test_generic_documents
Все DOCX/XLSX здесь собираются в памяти: внешние файлы не нужны.
"""

from __future__ import annotations

import unittest
from io import BytesIO

from docx import Document
from openpyxl import Workbook
from pypdf import PdfWriter
from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject

import app


class UploadedBytes:
    def __init__(self, name: str, data: bytes):
        self.name, self._data = name, data

    def getvalue(self) -> bytes:
        return self._data


def docx_file(name: str, *paragraphs: str) -> UploadedBytes:
    document = Document()
    for paragraph in paragraphs:
        document.add_paragraph(paragraph)
    buffer = BytesIO()
    document.save(buffer)
    return UploadedBytes(name, buffer.getvalue())


def xlsx_file(name: str, rows: list[tuple[str, str]]) -> UploadedBytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Функции"
    sheet.append(("Подразделение", "Обязанность"))
    for row in rows:
        sheet.append(row)
    buffer = BytesIO()
    workbook.save(buffer)
    return UploadedBytes(name, buffer.getvalue())


def pdf_file(name: str) -> UploadedBytes:
    """Две строки текста в PDF; библиотека pypdf уже требуется приложению."""
    writer = PdfWriter()
    page = writer.add_blank_page(width=595, height=842)
    font = DictionaryObject({
        NameObject("/Type"): NameObject("/Font"),
        NameObject("/Subtype"): NameObject("/Type1"),
        NameObject("/BaseFont"): NameObject("/Helvetica"),
    })
    page[NameObject("/Resources")] = DictionaryObject({
        NameObject("/Font"): DictionaryObject({NameObject("/F1"): writer._add_object(font)}),
    })
    stream = DecodedStreamObject()
    stream.set_data(
        b"BT /F1 12 Tf 40 700 Td "
        b"(Operations team reviews supplier invoices each month) Tj "
        b"0 -20 Td "
        b"(Quality team verifies the monthly invoice review procedure) Tj ET"
    )
    page[NameObject("/Contents")] = writer._add_object(stream)
    buffer = BytesIO()
    writer.write(buffer)
    return UploadedBytes(name, buffer.getvalue())


class GenericDocumentsTests(unittest.TestCase):
    def test_structure_alone_does_not_prove_zero_function_risks(self) -> None:
        self.assertFalse(app.audit_ready({"units": [{"status": "Создано"}], "mappings": [],
                                          "conclusion": "Создано одно подразделение."}, rejected=0))

    def test_unnumbered_word_document_has_verifiable_paragraph_sources(self) -> None:
        documents = app.load_documents([docx_file(
            "описание_функций.docx",
            "Отдел развития отвечает за подготовку плана внедрения новых услуг для клиентов компании.",
            "Служба качества отвечает за проверку выполнения плана внедрения новых услуг для клиентов компании.",
        )], "ПОСЛЕ")
        self.assertEqual(1, len(documents))
        document = documents[0]
        source = next(number for number, text in document.clauses.items()
                      if "подготовку плана внедрения новых услуг" in text)
        self.assertRegex(source, r"^Фрагмент\s+\d+$")
        verified, invalid = app.validated_refs([
            {"document": document.name, "clause": source,
             "quote": "Отдел развития отвечает за подготовку плана внедрения новых услуг"},
        ], {document.name: document}, "ПОСЛЕ")
        self.assertEqual(0, invalid)
        self.assertEqual([source], [item["clause"] for item in verified])
        self.assertIn(f"[{source}]", app.build_user_prompt(documents))
        self.assertNotIn("Казахтелеком", app.SYSTEM_PROMPT)

    def test_unnumbered_pdf_has_verifiable_page_and_line_sources(self) -> None:
        document = app.load_documents([pdf_file("roles_and_responsibilities.pdf")], "ДО")[0]
        source = "Страница 1, строка 2"
        self.assertIn(source, document.clauses)
        verified, invalid = app.validated_refs([
            {"document": document.name, "clause": source,
             "quote": "Quality team verifies the monthly invoice review procedure"},
        ], {document.name: document}, "ДО")
        self.assertEqual(0, invalid)
        self.assertEqual([source], [item["clause"] for item in verified])

    def test_new_and_removed_units_with_another_structure_heading(self) -> None:
        before = app.load_documents([docx_file(
            "структура_старой_компании.docx",
            "2.1. В организационную структуру входят следующие подразделения:",
            "а) Отдел сопровождения клиентов (ОСК).",
            "б) Служба развития продуктов (СРП).",
            "2.2. Эти подразделения выполняют разные задачи по обслуживанию клиентов и развитию продуктов.",
        )], "ДО")
        after = app.load_documents([docx_file(
            "новая_организационная_структура.docx",
            "2.1. В организационную структуру входят следующие подразделения:",
            "а) Отдел сопровождения клиентов (ОСК).",
            "б) Центр клиентского опыта (ЦКО).",
            "2.2. Эти подразделения выполняют разные задачи по обслуживанию клиентов и развитию продуктов.",
        )], "ПОСЛЕ")
        changes = app.compute_unit_changes(before, after)
        self.assertEqual(
            {("Создано", "ЦКО"), ("Упразднено", "СРП"), ("Сохранено без изменений", "ОСК")},
            {(change["status"], (change["unit_after"] if change["status"] != "Упразднено" else change["unit_before"]).split("(")[1].split(")")[0])
             for change in changes},
        )
        self.assertTrue(all(change["sources_before"] and change["sources_after"] for change in changes))

    def test_changed_duties_in_section_seven_are_reorganization(self) -> None:
        common = [
            "2.1. Компания состоит из следующих структурных подразделений:",
            "а) Отдел поддержки клиентов (ОПК).",
            "7.2. Начальник отдела поддержки клиентов:",
        ]
        before = app.load_documents([docx_file(
            "старые_обязанности.docx", *common,
            "7.2.1. Принимает запросы клиентов и фиксирует обращения в журнале.",
        )], "ДО")
        after = app.load_documents([docx_file(
            "новые_обязанности.docx", *common,
            "7.2.1. Принимает запросы клиентов и контролирует сроки решения обращений.",
        )], "ПОСЛЕ")
        changes = app.compute_unit_changes(before, after)
        self.assertEqual(1, len(changes))
        self.assertEqual("Реорганизовано", changes[0]["status"])
        self.assertTrue(any(ref["clause"] == "7.2.1" for ref in changes[0]["sources_before"]))
        self.assertTrue(any(ref["clause"] == "7.2.1" for ref in changes[0]["sources_after"]))

    def test_lost_hints_include_nonstandard_section(self) -> None:
        before = app.load_documents([docx_file(
            "процессы_до.docx",
            "7.2.1. Ответственный сотрудник еженедельно сверяет сроки выполнения заказов клиентов.",
            "7.2.2. Отдел обработки заявок регулярно принимает и регистрирует клиентские обращения.",
        )], "ДО")
        after = app.load_documents([docx_file(
            "процессы_после.docx",
            "7.2.1. Ответственный сотрудник передаёт заявки клиентов в службу поддержки пользователей.",
            "7.2.2. Отдел обработки заявок регулярно принимает и регистрирует клиентские обращения.",
        )], "ПОСЛЕ")
        self.assertIn("еженедельно сверяет сроки выполнения заказов", app.lost_candidate_hints(before + after))

    def test_overlap_hints_include_nonstandard_sections(self) -> None:
        after = app.load_documents([docx_file(
            "общие_обязанности.docx",
            "7.2.1. Руководитель отдела закупок выполняет контроль соблюдения сроков закупок и оценку выполнения договоров поставки.",
            "8.1.3. Руководитель отдела качества выполняет контроль соблюдения сроков закупок и оценку выполнения договоров поставки.",
        )], "ПОСЛЕ")
        hints = app.overlap_candidate_hints(after)
        self.assertIn("7.2.1", hints)
        self.assertIn("8.1.3", hints)

    def test_excel_from_unrelated_company_keeps_exact_sheet_and_row_sources(self) -> None:
        before = app.load_documents([xlsx_file("компания_до.xlsx", [
            ("Группа заказов", "Контроль выполнения заказов и соблюдения сроков поставки"),
        ])], "ДО")
        after = app.load_documents([xlsx_file("компания_после.xlsx", [
            ("Отдел логистики", "Контроль выполнения заказов и соблюдения сроков поставки"),
            ("Служба сопровождения", "Контроль выполнения заказов и соблюдения сроков поставки"),
        ])], "ПОСЛЕ")
        self.assertIn("Лист «Функции», строка 2", before[0].clauses)
        row = {
            "status": "Дублируется",
            "function_before": "Контроль выполнения заказов",
            "unit_before": "Группа заказов",
            "function_after": "Контроль выполнения заказов",
            "unit_after": "Отдел логистики и служба сопровождения",
            "sources_before": [{"document": before[0].name, "clause": "Лист «Функции», строка 2", "quote": "Контроль выполнения заказов и соблюдения сроков поставки"}],
            "sources_after": [{"document": after[0].name, "clause": f"Лист «Функции», строка {number}", "quote": "Контроль выполнения заказов и соблюдения сроков поставки"} for number in (2, 3)],
        }
        report, rejected = app.validate_report({"units": [], "mappings": [row]}, before + after)
        self.assertEqual(0, rejected)
        self.assertEqual(1, len(report["mappings"]))
        self.assertTrue(app.audit_ready(report, rejected))

    def test_fallback_must_reject_false_abolition_when_unit_still_listed(self) -> None:
        before = app.load_documents([docx_file(
            "старое_положение.docx",
            "2.1. В организационную структуру входят:",
            "а) Отдел снабжения и логистики (ОСЛ).",
            "2.2. Отдел снабжения и логистики отвечает за закупки, снабжение и доставку.",
        )], "ДО")
        after = app.load_documents([docx_file(
            "новое_положение.docx",
            "2.1. В организационную структуру входят:",
            "а) Отдел снабжения и логистики (ОСЛ).",
            "2.2. Отдел снабжения и логистики отвечает за закупки, снабжение и доставку.",
        )], "ПОСЛЕ")
        has_structural_comparison = bool(app.compute_unit_changes(before, after))
        old_ref = {"document": before[0].name, "clause": "2.1", "quote": "Отдел снабжения и логистики (ОСЛ)"}
        new_ref = {"document": after[0].name, "clause": "2.1", "quote": "Отдел снабжения и логистики (ОСЛ)"}
        raw = {"units": [{"status": "Упразднено", "unit_before": "Отдел снабжения и логистики (ОСЛ)",
                          "unit_after": "—", "rationale": "Подразделение якобы упразднено", "sources_before": [old_ref], "sources_after": [new_ref]}],
               "mappings": []}
        report, rejected = app.validate_report(raw, before + after)
        self.assertFalse(any(row["status"] == "Упразднено" for row in report["units"]))
        if not has_structural_comparison:
            self.assertGreater(rejected, 0)

    def test_fallback_rejects_false_abolition_without_known_list_heading(self) -> None:
        source = "Отдел снабжения и логистики (ОСЛ) отвечает за закупки и доставку оборудования клиентам."
        before = app.load_documents([docx_file("до.docx", source)], "ДО")
        after = app.load_documents([docx_file("после.docx", source)], "ПОСЛЕ")
        self.assertFalse(app.compute_unit_changes(before, after))
        row = {"status": "Упразднено", "unit_before": "Отдел снабжения и логистики (ОСЛ)", "unit_after": "—",
               "rationale": "Ошибочное утверждение об упразднении",
               "sources_before": [{"document": before[0].name, "clause": "Фрагмент 1", "quote": source}],
               "sources_after": [{"document": after[0].name, "clause": "Фрагмент 1", "quote": source}]}
        report, rejected = app.validate_report({"units": [row], "mappings": []}, before + after)
        self.assertGreater(rejected, 0)
        self.assertFalse(report["units"])

    def test_second_pass_duplicate_then_another_function_does_not_crash(self) -> None:
        before = app.load_documents([docx_file("до.docx",
            "Команда планирования выполняет контроль выполнения графика заказа клиентов.",
            "Группа кадров готовит график обучения сотрудников и определяет темы курса.")], "ДО")
        after = app.load_documents([docx_file("после.docx",
            "Группа кадров готовит график обучения сотрудников и определяет темы курса.",
            "Команда поддержки организует распределение новых заказов между консультантами.")], "ПОСЛЕ")

        def ref(doc, number, quote):
            return {"document": doc.name, "clause": f"Фрагмент {number}", "quote": quote}

        risk = {"status": "Утрачена", "function_before": "Контроль графика заказа", "unit_before": "Команда планирования",
                "function_after": "Сохранение контроля не обнаружено", "unit_after": "Команда поддержки",
                "sources_before": [ref(before[0], 1, "контроль выполнения графика заказа клиентов")],
                "sources_after": [ref(after[0], 2, "распределение новых заказов между консультантами")]}
        retained = {"status": "Сохранена", "function_before": "График обучения", "unit_before": "Группа кадров",
                    "function_after": "График обучения", "unit_after": "Группа кадров",
                    "sources_before": [ref(before[0], 2, "готовит график обучения сотрудников и определяет темы курса")],
                    "sources_after": [ref(after[0], 1, "готовит график обучения сотрудников и определяет темы курса")]}
        report, rejected = app.validate_report({"units": [], "mappings": [risk, risk, retained]}, before + after)
        self.assertEqual(0, rejected)
        self.assertEqual(["Утрачена", "Сохранена"], [row["status"] for row in report["mappings"]])


if __name__ == "__main__":
    unittest.main()
