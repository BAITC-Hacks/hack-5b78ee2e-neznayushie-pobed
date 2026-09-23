"""Контрольные случаи HackAlem по обезличенным редакциям положения №8 и №9.

Запуск после установки requirements.txt: python -m unittest -v test_audit_acceptance
Файлы-фикстуры: upload/Положение_..._редакция_{8,9}_обезличено.docx.
Если файлы лежат в другом месте, задайте TEST_BEFORE_DOCX и TEST_AFTER_DOCX.
"""

from __future__ import annotations

import os
import re
import unittest
from io import BytesIO
from pathlib import Path

from openpyxl import Workbook, load_workbook

import app


ROOT = Path(__file__).resolve().parent


class UploadedDocument:
    """Минимальный интерфейс файла, полученного Streamlit file_uploader."""

    def __init__(self, path: Path):
        self.name = path.name
        self._data = path.read_bytes()

    def getvalue(self) -> bytes:
        return self._data


class UploadedBytes:
    def __init__(self, name: str, data: bytes):
        self.name, self._data = name, data

    def getvalue(self) -> bytes:
        return self._data


def source_path(version: int) -> Path:
    override = os.getenv(f"TEST_{'BEFORE' if version == 8 else 'AFTER'}_DOCX")
    path = Path(override) if override else (
        ROOT / "upload" / f"Положение_о_внутреннем_аудите_редакция_{version}_обезличено.docx"
    )
    if not path.is_file():
        raise FileNotFoundError(
            f"Для контрольных тестов нужен {path}. "
            "Задайте TEST_BEFORE_DOCX и TEST_AFTER_DOCX, если документы находятся в другой папке."
        )
    return path


def unit_mentions(item: dict, abbreviation: str) -> bool:
    """Ищем подразделение, не принимая совпадения в объяснении или ссылках."""
    fields = " ".join((item.get("unit_before") or "", item.get("unit_after") or ""))
    return re.search(r"(?<!\w)" + re.escape(abbreviation) + r"(?!\w)", fields, re.I) is not None


class RealDocumentsAcceptanceTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls) -> None:
        cls.before = app.load_documents([UploadedDocument(source_path(8))], "ДО")
        cls.after = app.load_documents([UploadedDocument(source_path(9))], "ПОСЛЕ")
        cls.before_doc = cls.before[0]
        cls.after_doc = cls.after[0]

    def test_source_clause_3_4_is_extracted_with_actual_units(self) -> None:
        before = self.before_doc.clauses["3.4"]
        after = self.after_doc.clauses["3.4"]
        self.assertIn("БВА состоит из следующих структурных подразделений", before)
        self.assertIn("БВА состоит из следующих структурных подразделений", after)
        self.assertIn("(ДНМ)", before)
        self.assertIn("(ДККМ)", before)
        self.assertNotIn("(ДИТААД)", before)
        self.assertNotIn("(ДОА)", before)
        for abbreviation in ("ДИТААД", "ДОА", "ДНМ", "ДККМ"):
            self.assertIn(f"({abbreviation})", after)

    def test_unit_change_baseline_preserves_parent_and_identifies_new_units(self) -> None:
        changes = app.compute_unit_changes(self.before, self.after)
        self.assertIsInstance(changes, list)
        self.assertTrue(changes, "Реальные редакции дают как минимум два созданных департамента")

        for abbreviation in ("ДИТААД", "ДОА"):
            matches = [item for item in changes if unit_mentions(item, abbreviation)]
            self.assertEqual(1, len(matches), f"Структурная единица {abbreviation} должна встретиться ровно раз")
            self.assertEqual("Создано", matches[0]["status"])
            self.assertTrue(matches[0].get("sources_after"), "Нужна ссылка на п.3.4 редакции №9")

        for abbreviation in ("ДНМ", "ДККМ"):
            matches = [item for item in changes if unit_mentions(item, abbreviation)]
            self.assertEqual(1, len(matches), f"Существующий департамент {abbreviation} нельзя терять")
            self.assertEqual("Реорганизовано", matches[0]["status"])
            self.assertTrue(matches[0].get("sources_before"))
            self.assertTrue(matches[0].get("sources_after"))

        self.assertFalse(
            any(item["status"] == "Упразднено" and unit_mentions(item, "БВА") for item in changes),
            "П.3.4 обеих редакций сохраняет БВА как родительский блок",
        )
        self.assertFalse(
            any("Директор направления внутреннего аудита" in (item.get("unit_before") or "") for item in changes),
            "Должность директора в п.3.5 №8 нельзя выдавать за упраздненный департамент",
        )

    def test_existing_clause_numbers_do_not_validate_a_false_abolition(self) -> None:
        before_ref = {"document": self.before_doc.name, "clause": "3.4"}
        after_ref = {"document": self.after_doc.name, "clause": "3.4"}
        raw = {
            "units": [{
                "status": "Упразднено",
                "unit_before": "БВА",
                "unit_after": "",
                "rationale": "Блок внутреннего аудита упразднен",
                "sources_before": [before_ref],
                "sources_after": [after_ref],
            }],
            "mappings": [],
            "conclusion": "",
            "conclusion_sources": [],
            "recommendations": [],
        }
        report, _ = app.validate_report(raw, self.before + self.after)
        self.assertFalse(
            any(item["status"] == "Упразднено" and unit_mentions(item, "БВА") for item in report["units"]),
            "Оба номера 3.4 существуют, но их текст опровергает вывод об упразднении БВА",
        )

    def test_lost_explicit_deadline_control_and_real_overlap_have_source_text(self) -> None:
        before_deadlines = self.before_doc.clauses["5.3.4"]
        after_team = self.after_doc.clauses["5.3.4"] + "\n" + self.after_doc.clauses["5.3.5"]
        self.assertIn("контроль сроков выполнения графика", before_deadlines)
        self.assertIn("формирование графика", after_team)
        self.assertNotIn("контроль сроков выполнения графика", after_team)

        after_audit = self.after_doc.clauses["5.3.8"]
        after_monitoring = self.after_doc.clauses["5.4.5"]
        self.assertIn("анализируют результаты проверок БВА и непрерывного аудита", after_audit)
        self.assertIn("анализирует результаты непрерывного аудита", after_monitoring)
        self.assertIn("ДИТААД", self.after_doc.clauses["5.3.2"])
        self.assertIn("ДОА", self.after_doc.clauses["5.3.2"])
        self.assertIn("контроль сроков выполнения графика", app.lost_candidate_hints(self.before + self.after))
        self.assertTrue(any(
            "5.3.8" in line and "5.4.5" in line
            for line in app.overlap_candidate_hints(self.before + self.after).splitlines()
        ))


class IncompleteReportTests(unittest.TestCase):
    def test_same_headcount_with_changed_duties_is_reorganized(self) -> None:
        structure = "3.4. БВА состоит из следующих структурных подразделений:\nа. Департамент непрерывного мониторинга (ДНМ)."
        staff = "3.5. Директору ДНМ подчиняются работники ДНМ в составе должностей: аудитор."
        heading = "5.4. Директор департамента непрерывного мониторинга:"
        before = app.SourceDocument("ДО / структура.docx", "ДО", "", {"3.4": structure, "3.5": staff,
            "5.4": heading, "5.4.1": "5.4.1. Проверяет график выполнения."})
        after = app.SourceDocument("ПОСЛЕ / структура.docx", "ПОСЛЕ", "", {"3.4": structure, "3.5": staff,
            "5.4": heading, "5.4.1": "5.4.1. Проверяет качество отчётов."})
        changes = app.compute_unit_changes([before], [after])
        self.assertEqual("Реорганизовано", changes[0]["status"])
        self.assertEqual("5.4.1", changes[0]["sources_before"][-1]["clause"])
        self.assertEqual("5.4.1", changes[0]["sources_after"][-1]["clause"])

    def test_rejected_findings_make_zero_metrics_unreliable(self) -> None:
        report = {
            "units": [{"status": "Создано", "unit_after": "ДОА"}],
            "mappings": [{"status": "Сохранена"}],
            "conclusion": "Вывод по доступным строкам",
        }
        # Сценарий пользователя: из нескольких найденных строк осталось только
        # сохранение функций, остальные отброшены из-за неверных ссылок.
        self.assertFalse(app.audit_ready(report, rejected=5))

    def test_empty_report_is_not_ready_for_export(self) -> None:
        self.assertFalse(app.audit_ready({"units": [], "mappings": [], "conclusion": ""}, rejected=0))

    def test_complete_report_can_be_exported(self) -> None:
        report = {"units": [{"status": "Создано"}], "mappings": [{"status": "Дублируется"}], "conclusion": "Вывод"}
        self.assertTrue(app.audit_ready(report, rejected=0))


class EvidenceValidationTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls) -> None:
        cls.before = app.load_documents([UploadedDocument(source_path(8))], "ДО")
        cls.after = app.load_documents([UploadedDocument(source_path(9))], "ПОСЛЕ")
        cls.before_doc, cls.after_doc = cls.before[0], cls.after[0]

    def ref(self, doc, clause, quote):
        return {"document": doc.name, "clause": clause, "quote": quote}

    def test_three_real_risks_are_verified_with_quotes(self) -> None:
        before, after = self.before_doc, self.after_doc
        old_deadline = self.ref(before, "5.3.4(б)", "контроль сроков выполнения графика по проверке проектной командой")
        new_schedule = self.ref(after, "5.3.4(в)", "формирование графика и постановку задач")
        new_goals = self.ref(after, "5.3.5(а)", "контроль выполнения целей аудита")
        old_analytics = self.ref(before, "5.3.7", "анализирует результаты проверок БВА и непрерывного аудита")
        new_analytics = self.ref(after, "5.3.8", "анализируют результаты проверок БВА и непрерывного аудита")
        new_monitoring = self.ref(after, "5.4.5", "анализирует результаты непрерывного аудита")
        rows = [
            {"status": "Утрачена", "function_before": "Прямой контроль сроков графика проверки", "unit_before": "Директор направления", "function_after": "Прямое закрепление не найдено", "unit_after": "ДИТААД/ДОА", "sources_before": [old_deadline], "sources_after": [new_schedule, new_goals]},
            {"status": "Дублируется", "function_before": "Анализ результатов непрерывного аудита", "unit_before": "Направление аудита", "function_after": "Анализ результатов непрерывного аудита", "unit_after": "ДИТААД/ДОА и ДНМ", "sources_before": [old_analytics], "sources_after": [new_analytics, new_monitoring]},
            {"status": "Конфликт", "function_before": "Проверка и контроль команды", "unit_before": "Директор направления", "function_after": "Проверка и контроль качества своей проектной команды", "unit_after": "ДИТААД/ДОА", "sources_before": [self.ref(before, "5.3.4", "организация контроля качества работы проектной команды")], "sources_after": [self.ref(after, "5.3.5", "проводят проверки и обеспечивают выполнение плана работ БВА"), self.ref(after, "5.3.5(б)", "контроль качества работы проектной команды")]},
        ]
        report, rejected = app.validate_report({"units": [], "mappings": rows}, self.before + self.after)
        self.assertEqual(0, rejected)
        self.assertEqual({"Утрачена", "Дублируется", "Конфликт"}, {row["status"] for row in report["mappings"]})
        self.assertTrue(app.audit_ready(report, rejected))
        self.assertIn("создано подразделений: 2", report["conclusion"])
        workbook = load_workbook(BytesIO(app.export_excel(app.mapping_frame(report), report)))
        self.assertEqual(["Сопоставление", "Подразделения", "Заключение", "Источники"], workbook.sheetnames)
        self.assertEqual(4, workbook["Сопоставление"].max_row)
        self.assertTrue(any(
            "контроль сроков" in str(row[2].value).casefold()
            for row in list(workbook["Источники"].rows)[1:]
        ))

    def test_existing_number_with_false_quote_is_not_accepted(self) -> None:
        bad = self.ref(self.after_doc, "3.4", "БВА полностью упразднён")
        verified, invalid = app.validated_refs([bad], {self.after_doc.name: self.after_doc}, "ПОСЛЕ")
        self.assertFalse(verified)
        self.assertEqual(1, invalid)

    def test_quote_must_belong_to_the_cited_subpoint(self) -> None:
        quote = "контроль сроков выполнения графика по проверке проектной командой"
        bad = self.ref(self.before_doc, "5.3.4(а)", quote)  # Цитата находится в б, а не в а.
        good = self.ref(self.before_doc, "5.3.4(б)", quote)
        verified, invalid = app.validated_refs([bad, good], {self.before_doc.name: self.before_doc}, "ДО")
        self.assertEqual(1, invalid)
        self.assertEqual(["5.3.4(б)"], [item["clause"] for item in verified])

    def test_second_pass_does_not_erase_or_duplicate_first_pass_risk(self) -> None:
        row = {
            "status": "Утрачена", "function_before": "Контроль сроков выполнения графика", "unit_before": "Директор",
            "function_after": "Прямое закрепление не найдено", "unit_after": "Директора",
            "sources_before": [self.ref(self.before_doc, "5.3.4(б)", "контроль сроков выполнения графика по проверке проектной командой")],
            "sources_after": [self.ref(self.after_doc, "5.3.4(в)", "формирование графика и постановку задач")],
        }
        general = {"units": [], "mappings": [row]}
        merged, unexpected = app.combine_api_results(general, {"mappings": []})
        self.assertFalse(unexpected)
        report, rejected = app.validate_report(merged, self.before + self.after)
        self.assertEqual(0, rejected)
        self.assertEqual(1, sum(item["status"] == "Утрачена" for item in report["mappings"]))

        combined, _ = app.combine_api_results(general, {"mappings": [row]})
        report, rejected = app.validate_report(combined, self.before + self.after)
        self.assertEqual(0, rejected)
        self.assertEqual(1, sum(item["status"] == "Утрачена" for item in report["mappings"]))

    def test_excel_sheet_row_is_a_valid_source(self) -> None:
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Структура"
        sheet["A2"] = "Департамент операционного аудита"
        buffer = BytesIO()
        workbook.save(buffer)
        document = app.load_documents([UploadedBytes("структура.xlsx", buffer.getvalue())], "ПОСЛЕ")[0]
        anchor = "Лист «Структура», строка 2"
        self.assertIn(anchor, document.clauses)
        refs, invalid = app.validated_refs(
            [{"document": document.name, "clause": anchor, "quote": "Департамент операционного аудита"}],
            {document.name: document}, "ПОСЛЕ",
        )
        self.assertEqual(0, invalid)
        self.assertEqual(anchor, refs[0]["number"])


if __name__ == "__main__":
    unittest.main()
