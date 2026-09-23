"""Чтение входных Excel-файлов в текст со ссылками на исходные строки.

Формулы .xlsx передаются как текст формулы и никогда не выполняются. Библиотека
xlrd возвращает сохранённые значения ячеек .xls, включая результаты формул.
Пустые строки пропускаются, но исходные номера непустых строк сохраняются.
"""

from __future__ import annotations

from datetime import date, datetime, time
from io import BytesIO
from math import isfinite
from typing import Any, Iterator
from zipfile import ZipFile

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
import xlrd


MAX_SHEETS = 100
MAX_ROWS_PER_SHEET = 25_000
MAX_COLUMNS_PER_SHEET = 1_024
MAX_NONEMPTY_CELLS = 100_000
MAX_CELL_CHARACTERS = 50_000
MAX_TEXT_CHARACTERS = 4_000_000
MAX_UNCOMPRESSED_XLSX_BYTES = 64 * 1024 * 1024


def _cell_text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, datetime):
        return value.isoformat(sep=" ")
    if isinstance(value, (date, time)):
        return value.isoformat()
    if isinstance(value, float) and not isfinite(value):
        return str(value)
    return str(value).replace("\x00", "").replace("\xa0", " ").replace("\r", " / ").replace("\n", " / ").strip()


def _xlsx_sheets(data: bytes) -> Iterator[tuple[str, int, int, Iterator[tuple[int, Iterator[tuple[int, Any]]]]]]:
    # ZIP-книга может быть маленькой при загрузке, но гигантской после распаковки.
    with ZipFile(BytesIO(data)) as archive:
        if sum(member.file_size for member in archive.infolist()) > MAX_UNCOMPRESSED_XLSX_BYTES:
            raise ValueError("Excel-книга слишком велика после распаковки; разделите её на части.")
    workbook = load_workbook(BytesIO(data), read_only=True, data_only=False, keep_links=False)
    try:
        for sheet in workbook.worksheets:
            def rows(sheet=sheet):
                for number, cells in enumerate(sheet.iter_rows(), 1):
                    yield number, ((column, cell.value) for column, cell in enumerate(cells, 1))

            yield sheet.title, sheet.max_row or 0, sheet.max_column or 0, rows()
    finally:
        workbook.close()


def _xls_sheets(data: bytes) -> Iterator[tuple[str, int, int, Iterator[tuple[int, Iterator[tuple[int, Any]]]]]]:
    workbook = xlrd.open_workbook(file_contents=data, on_demand=True)
    try:
        for sheet in workbook.sheets():
            def rows(sheet=sheet):
                for index in range(sheet.nrows):
                    def cells(index=index):
                        for col, cell in enumerate(sheet.row(index), 1):
                            if cell.ctype in (xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK):
                                continue
                            if cell.ctype == xlrd.XL_CELL_DATE:
                                value = xlrd.xldate.xldate_as_datetime(cell.value, workbook.datemode)
                            elif cell.ctype == xlrd.XL_CELL_ERROR:
                                value = f"Ошибка Excel {xlrd.error_text_from_code.get(cell.value, cell.value)}"
                            else:
                                value = cell.value
                            yield col, value

                    yield index + 1, cells()

            yield sheet.name, sheet.nrows, sheet.ncols, rows()
    finally:
        workbook.release_resources()


def extract_excel(data: bytes, extension: str) -> tuple[str, dict[str, str]]:
    """Вернуть ``(текст, {якорь строки: выдержка})`` для .xlsx или .xls.

    Например, ключ ``Лист «Функции», строка 7`` сохраняет имя листа и реальный
    номер строки Excel. В выдержке ячейки помечаются буквами столбцов.
    Превышение ограничений вызывает ошибку вместо скрытой обрезки документа.
    """
    kind = extension.lower().lstrip(".")
    if kind not in {"xlsx", "xls"}:
        raise ValueError("Поддерживаются форматы .xlsx и .xls.")
    sheets = _xlsx_sheets(data) if kind == "xlsx" else _xls_sheets(data)
    anchors: dict[str, str] = {}
    lines: list[str] = []
    nonempty_cells = 0
    total_characters = 0
    sheet_count = 0

    for name, row_count, column_count, rows in sheets:
        sheet_count += 1
        if sheet_count > MAX_SHEETS:
            raise ValueError(f"В книге больше {MAX_SHEETS} листов; разделите её на части.")
        if row_count > MAX_ROWS_PER_SHEET or column_count > MAX_COLUMNS_PER_SHEET:
            raise ValueError(
                f"Лист «{name}» слишком большой ({row_count} строк, {column_count} столбцов); "
                "разделите книгу на тематические части."
            )

        for row_number, cells in rows:
            if row_number > MAX_ROWS_PER_SHEET:
                raise ValueError(f"В листе «{name}» более {MAX_ROWS_PER_SHEET} строк.")
            columns: list[str] = []
            for column_number, value in cells:
                cell = _cell_text(value)
                if not cell:
                    continue
                if len(cell) > MAX_CELL_CHARACTERS:
                    raise ValueError(f"Слишком длинная ячейка в листе «{name}», строке {row_number}.")
                nonempty_cells += 1
                if nonempty_cells > MAX_NONEMPTY_CELLS:
                    raise ValueError("В книге слишком много заполненных ячеек; разделите её на части.")
                columns.append(f"{get_column_letter(column_number)}: {cell}")
            if not columns:
                continue
            anchor = f"Лист «{name}», строка {row_number}"
            excerpt = " | ".join(columns)
            line = f"[{anchor}] {excerpt}"
            total_characters += len(line) + 1
            if total_characters > MAX_TEXT_CHARACTERS:
                raise ValueError("Текст книги слишком большой для аудита; разделите её на части.")
            anchors[anchor] = excerpt
            lines.append(line)

    if not anchors:
        raise ValueError("В Excel-файле нет заполненных ячеек.")
    return "\n".join(lines), anchors
